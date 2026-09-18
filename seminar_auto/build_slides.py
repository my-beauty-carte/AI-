import base64
import tempfile
from pathlib import Path

from openai import OpenAI
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.text.text import _Run
from pptx.util import Inches, Pt

from . import config

# デフォルトのpptxテーマは東アジア用フォントが未指定で、日本語用のフォールバックが
# Windows専用の「ＭＳ Ｐゴシック」になっている。実行環境を問わず表示できるよう、
# 各テキストに東アジア用フォントを明示的に指定する。
#
# macOS標準の「Hiragino Sans」はシステム保護領域のフォントコンテナに格納されており、
# LibreOfficeをheadless(バックグラウンド)モードで実行した際にはこれを見つけられず、
# 日本語部分だけ表示されなくなることを確認した。通常のフォントファイルとして
# インストールされる「Noto Sans JP」(`brew install --cask font-noto-sans-jp`)は
# headlessモードでも問題なく認識されるため、こちらを使う。
JAPANESE_FONT = "Noto Sans JP"

IMAGE_PROMPT_TEMPLATE = (
    "シンプルなフラットデザインのアイコン風イラスト。テーマ: {title}。"
    "ミニマルな配色で、背景はシンプルな単色。文字やテキストは一切含めないこと。"
)


def _set_japanese_font(run: _Run, font_name: str) -> None:
    run.font.name = font_name
    rPr = run.font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = OxmlElement("a:ea")
        rPr.append(ea)
    ea.set("typeface", font_name)


def _generate_slide_image(title: str) -> bytes:
    client = OpenAI()
    response = client.images.generate(
        model=config.OPENAI_IMAGE_MODEL,
        prompt=IMAGE_PROMPT_TEMPLATE.format(title=title),
        size="1024x1024",
        n=1,
        response_format="b64_json",
    )
    return base64.b64decode(response.data[0].b64_json)


def build(slides: list[dict], output_path: Path) -> None:
    prs = Presentation()
    layout = prs.slide_layouts[1]  # タイトル + 本文
    for slide_data in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = slide_data["title"]
        for run in slide.shapes.title.text_frame.paragraphs[0].runs:
            _set_japanese_font(run, JAPANESE_FONT)

        # 箇条書きは左半分、イラストは右半分に配置する。
        body = slide.placeholders[1]
        body.left = Inches(0.5)
        body.top = Inches(1.7)
        body.width = Inches(5.3)
        body.height = Inches(5.3)
        text_frame = body.text_frame
        text_frame.clear()
        for i, bullet in enumerate(slide_data["bullets"]):
            paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(28)
            for run in paragraph.runs:
                _set_japanese_font(run, JAPANESE_FONT)

        image_bytes = _generate_slide_image(slide_data["title"])
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
            f.write(image_bytes)
            image_path = Path(f.name)
        try:
            slide.shapes.add_picture(str(image_path), Inches(6.0), Inches(1.9), height=Inches(3.6))
        finally:
            image_path.unlink()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
